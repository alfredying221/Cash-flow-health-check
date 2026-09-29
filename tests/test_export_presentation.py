from io import BytesIO
import unittest

from openpyxl import load_workbook

from qa_gate1_regression import CASES, load_case, scenario_rows
from senalo_analysis.exports import export_excel
from fulfilment.analysis_processor import default_forecast_assumptions


class ExportPresentationTests(unittest.TestCase):
    def test_health_scores_are_points_and_assumptions_remain_percentages(self):
        for case in CASES:
            with self.subTest(fixture=case["file"]):
                history, metrics, score, label, breakdown = load_case(
                    case["file"], case["opening_cash"]
                )
                forecast, scenarios, details = scenario_rows(history, metrics)
                assumptions = {
                    "Business Type": case["business_type"],
                    **default_forecast_assumptions(metrics),
                    "Downside Sales Adjustment": -0.15,
                    "Upside Sales Adjustment": 0.15,
                }
                data = export_excel(
                    history, forecast, scenarios, breakdown, assumptions,
                    case["business_type"], details, metrics, label,
                )
                workbook = load_workbook(BytesIO(data))
                sheet = workbook["Health Score"]
                for column in ("Score", "Maximum Score"):
                    index = [cell.value for cell in sheet[1]].index(column) + 1
                    cells = list(sheet.iter_cols(min_col=index, max_col=index, min_row=2))[0]
                    self.assertEqual([cell.value for cell in cells], breakdown[column].tolist())
                    self.assertTrue(all(cell.number_format == "0" for cell in cells))
                self.assertEqual(sum(row[1].value for row in list(sheet.rows)[1:]), score)
                for name, value in list(workbook["Assumptions"].rows)[2:]:
                    self.assertEqual(value.number_format, "0.0%", name.value)
                    self.assertEqual(value.value, assumptions[name.value])


if __name__ == "__main__":
    unittest.main()
